import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import httpx
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from websheet.core import Column, ExtractionError, extract, export_rows, fetch, origin, public_addresses, safe_cell
from websheet.store import Store
from websheet.app import create_app
from websheet.ai import suggest

FIELDS = [Column(name='名称', selector='h1'), Column(name='价格', selector='.price')]

class CoreTests(unittest.TestCase):
    def test_demo_30(self):
        for n in range(1, 31):
            result = extract(fetch(f'demo://catalog/{n}'), FIELDS)
            self.assertEqual(result['价格'], f'{n * 13}.00')

    def test_selectors(self):
        page = '<h1 id="name" class="title">Hello <b>world</b></h1><a href="/safe">Link</a>'
        for selector in ['h1', '#name', '.title', 'h1.title', 'h1#name']:
            self.assertEqual(extract(page, [Column(name='x', selector=selector)])['x'], 'Hello world')
        self.assertEqual(extract(page, [Column(name='link', selector='a', attribute='href')])['link'], '/safe')

    def test_stale_ambiguous_empty(self):
        for page in ['<p>no</p>', '<h1>a</h1><h1>b</h1>', '<h1></h1>']:
            with self.assertRaises(ExtractionError):
                extract(page, FIELDS[:1])

    def test_selector_injection(self):
        for selector in ['h1 | //script', '*', 'h1 p', '#x"]', ':contains(foo)']:
            with self.assertRaises(ValueError):
                Column(name='x', selector=selector)

    def test_duplicate_reserved_names(self):
        for fields in [[FIELDS[0], FIELDS[0]], [Column(name='source_id', selector='h1')]]:
            with self.assertRaises(ExtractionError):
                extract('<h1>x</h1>', fields)

    def test_url_restrictions(self):
        for url in ['http://example.com', 'file:///etc/passwd', 'https://a:444/', 'https://u:p@a/', 'https://a/#x', 'demo://catalog/0', 'demo://catalog/1?x', 'data:hi']:
            with self.assertRaises(ExtractionError):
                origin(url)

    def test_ssrf_private_dns_mixed(self):
        for address in ['127.0.0.1', '10.1.1.1', '169.254.169.254', '::1', '192.0.2.1', '::ffff:127.0.0.1']:
            with patch('socket.getaddrinfo', return_value=[(None, None, None, None, (address, 443))]):
                with self.assertRaises(ExtractionError):
                    public_addresses('example.org')
        with patch('socket.getaddrinfo', return_value=[(None, None, None, None, ('8.8.8.8', 443)), (None, None, None, None, ('127.0.0.1', 443))]):
            with self.assertRaises(ExtractionError):
                public_addresses('example.org')

    def test_public_dns(self):
        with patch('socket.getaddrinfo', return_value=[(None, None, None, None, ('8.8.8.8', 443))]):
            self.assertEqual(public_addresses('example.org'), ['8.8.8.8'])

    def test_formula_escape(self):
        for value in ['=1+1', '+SUM(A1)', '-1', '@SUM(A1)', '\tstuff', '  =cmd']:
            self.assertTrue(safe_cell(value).startswith("'"))
        self.assertEqual(safe_cell('normal'), 'normal')
        self.assertEqual(safe_cell('a\x01b'), 'ab')

    def test_network_headers_redirect_size(self):
        from unittest.mock import MagicMock
        conn = MagicMock()
        response = conn.getresponse.return_value
        response.status = 302
        with patch('websheet.core.public_addresses', return_value=['8.8.8.8']), patch('websheet.core.PinnedHTTPS', return_value=conn):
            with self.assertRaises(ExtractionError): fetch('https://example.org?x=1')
            self.assertEqual(conn.request.call_args.args[1], '/?x=1')
            conn.close.assert_called()
            response.status = 200
            response.getheader.side_effect = lambda key, default='': {'Content-Type': 'text/html', 'Content-Encoding': 'identity'}.get(key, default)
            response.read.return_value = b'a' * 1_000_001
            with self.assertRaises(ExtractionError): fetch('https://example.org/')

class JobTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.store = Store(Path(self.tmp.name) / 'jobs.sqlite3')
    def tearDown(self): self.tmp.cleanup()

    def test_batch_export_and_idempotence(self):
        job = self.store.create([f'demo://catalog/{n}' for n in range(1, 31)] + ['demo://catalog/1'], FIELDS, True)
        done = self.store.run(job['id'])
        self.assertEqual((done['total'], done['completed'], done['failed']), (30, 30, 0))
        again = self.store.run(job['id'], fetcher=lambda _: self.fail('must not refetch completed rows'))
        self.assertEqual(done, again)
        csv_data = export_rows(done, 'csv').decode('utf-8-sig')
        self.assertEqual(len(csv_data.splitlines()), 31)
        workbook = load_workbook(io.BytesIO(export_rows(done, 'xlsx')))
        self.assertEqual(workbook.active.max_row, 31)
        self.assertEqual(workbook.active.cell(2, 2).value, 'demo://catalog/1')

    def test_failure_and_resume_after_crash(self):
        job = self.store.create(['demo://catalog/1', 'demo://catalog/2', 'demo://catalog/3'], FIELDS, True)
        def fail_second(url):
            if url.endswith('/2'): raise KeyboardInterrupt()
            return fetch(url)
        with self.assertRaises(KeyboardInterrupt): self.store.run(job['id'], fail_second)
        restored = Store(self.store.path)
        self.assertEqual(restored.get(job['id'])['completed'], 1)
        def wrong_third(url):
            if url.endswith('/1'): self.fail('success must be immutable')
            return '<h1>Changed</h1>' if url.endswith('/3') else fetch(url)
        done = restored.run(job['id'], wrong_third)
        self.assertEqual((done['completed'], done['failed']), (2, 1))
        self.assertIn('匹配 0', done['rows'][2]['error'])
        self.assertEqual(done['rows'][2]['data'], {})

    def test_review_origin_engine(self):
        for urls, reviewed, engine in [(['demo://catalog/1'], False, 'static'), (['https://a/x','https://b/x'],True,'static'), (['https://a/x'],True,'browser-demo')]:
            with self.assertRaises(ExtractionError): self.store.create(urls,FIELDS,reviewed,engine)

    def test_xlsx_controls_formula_headers(self):
        job = self.store.create(['demo://catalog/1'], [Column(name='=danger', selector='h1')], True)
        done = self.store.run(job['id'], lambda _: '<h1>\x01=bad</h1>')
        sheet = load_workbook(io.BytesIO(export_rows(done,'xlsx'))).active
        self.assertEqual(sheet.cell(1,5).value, "'=danger")
        self.assertEqual(sheet.cell(2,5).data_type, 's')
        self.assertEqual(sheet.cell(2,5).value, "'=bad")

class ApiTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.client = TestClient(create_app(Path(self.tmp.name)/'app.sqlite3'), base_url='http://127.0.0.1')
    def tearDown(self): self.tmp.cleanup()
    def test_end_to_end_demo(self):
        demo = self.client.get('/api/demo').json()
        preview = self.client.post('/api/preview', json={'url': demo['sample_url'], 'fields': demo['fields']})
        self.assertEqual(preview.status_code, 200)
        job = self.client.post('/api/jobs', json={**demo, 'reviewed':True}).json()
        done = self.client.post(f'/api/jobs/{job["id"]}/run').json()
        self.assertEqual(done['completed'], 30)
        for fmt in ['csv','xlsx']:
            export = self.client.get(f'/api/jobs/{job["id"]}/export?format={fmt}')
            self.assertEqual(export.status_code,200)
            self.assertGreater(len(export.content),100)
    def test_request_boundaries(self):
        self.assertEqual(self.client.post('/api/jobs', json={'urls':['demo://catalog/1'],'fields':[f.model_dump() for f in FIELDS]}).status_code,400)
        self.assertEqual(self.client.get('/api/demo', headers={'host':'evil.example'}).status_code,403)
        self.assertEqual(self.client.post('/api/preview', headers={'origin':'https://evil.example'},json={}).status_code,403)
        self.assertEqual(self.client.get('/api/jobs/missing').status_code,404)
        self.assertEqual(self.client.post('/api/suggest',json={'url':'demo://catalog/1','names':['名称']}).status_code,400)
    def test_no_secret_serialized(self):
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY':'secret-not-for-browser'}):
            self.assertNotIn('secret-not-for-browser', self.client.get('/api/demo').text)

class AiTests(unittest.TestCase):
    def test_mock_structured_model(self):
        def handler(request):
            self.assertEqual(request.url.host, 'api.deepseek.com')
            self.assertEqual(request.headers['Authorization'], 'Bearer test-key')
            return httpx.Response(200,json={'choices':[{'message':{'content':json.dumps({'fields':[FIELDS[0].model_dump()]})}}]})
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY':'test-key'}), httpx.Client(transport=httpx.MockTransport(handler)) as client:
            fields = suggest('<h1>Demo</h1><script>evil()</script>', ['名称'], client=client)
            self.assertEqual(fields[0].selector, 'h1')
    def test_no_key(self):
        with patch.dict(os.environ, {}, clear=True):
            with self.assertRaises(ExtractionError): suggest('<h1>x</h1>', ['名称'])
    def test_invalid_ai_output_not_executed(self):
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY':'test'}), httpx.Client(transport=httpx.MockTransport(lambda _: httpx.Response(200,json={'choices':[{'message':{'content':'{"fields":[{"name":"名称","selector":"//script"}]}'}}]}))) as client:
            with self.assertRaises(ExtractionError): suggest('<h1>x</h1>', ['名称'], client=client)

if __name__ == '__main__': unittest.main()
