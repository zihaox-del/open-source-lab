"""Run the fully offline synthetic catalog workflow; never a real website claim."""
import io
from pathlib import Path
import sys
import tempfile
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from websheet.core import Column, export_rows
from websheet.store import Store

with tempfile.TemporaryDirectory() as directory:
    store = Store(Path(directory) / 'demo.sqlite3')
    job = store.create([f'demo://catalog/{i}' for i in range(1, 31)], [Column(name='名称', selector='h1'), Column(name='价格', selector='.price')], True)
    result = store.run(job['id'])
    workbook = load_workbook(io.BytesIO(export_rows(result, 'xlsx')))
    assert (result['completed'], result['failed'], workbook.active.max_row) == (30, 0, 31)
    print('SYNTHETIC OFFLINE DEMO: 30/30 successful, 0 failures, XLSX 30 data rows + header; no website or model API accessed.')
