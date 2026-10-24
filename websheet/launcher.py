"""Portable local launcher. No installations, global settings, or remote listeners."""
import argparse
import io
import json
import os
from pathlib import Path
import socket
import signal
import sys
import tempfile
import threading
import time
import urllib.request
import webbrowser


def reserve_socket(port):
    sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    try:
        # Windows SO_REUSEADDR can let a second process steal an occupied port.
        if os.name == 'nt':
            sock.setsockopt(socket.SOL_SOCKET, socket.SO_EXCLUSIVEADDRUSE, 1)
        else:
            # Permit clean restart despite TIME_WAIT; never enable SO_REUSEPORT.
            sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        sock.bind(('127.0.0.1', port))
    except BaseException:
        sock.close()
        raise
    return sock


def open_when_ready(server, url, stop):
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    while not stop.wait(.1):
        if not server.started:
            continue
        try:
            with opener.open(url + '/api/health', timeout=1) as response:
                result = json.load(response)
            if result == {'status': 'ok', 'app': 'websheet-agent'} and not stop.is_set():
                webbrowser.open(url)
                return
        except (OSError, ValueError):
            continue


def self_check():
    from openpyxl import load_workbook
    from .core import Column, export_rows
    from .store import Store
    with tempfile.TemporaryDirectory() as directory:
        store = Store(Path(directory) / 'selfcheck.sqlite3')
        job = store.create([f'demo://catalog/{i}' for i in range(1, 31)],
                           [Column(name='名称', selector='h1'), Column(name='价格', selector='.price')], True)
        result = store.run(job['id'])
        workbook = load_workbook(io.BytesIO(export_rows(result, 'xlsx')))
        if (result['completed'], result['failed'], workbook.active.max_row) != (30, 0, 31):
            raise RuntimeError('Offline self-check failed')
        if not (Path(__file__).parent / 'static' / 'index.html').is_file():
            raise RuntimeError('Bundled dashboard is missing')
        store.shutdown()
    print('OFFLINE SELF-CHECK OK: 30/30 synthetic rows and XLSX; no network or API key.')
    return 0


def main(argv=None):
    # Redirected Windows consoles may use a legacy code page. A diagnostic
    # character must never turn successful startup into an encoding failure.
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(errors='replace')
    parser = argparse.ArgumentParser(description='网页搬表 · WebSheet Agent（仅本机）')
    parser.add_argument('--no-browser', action='store_true', help='不自动打开浏览器')
    parser.add_argument('--port', type=int, default=8765)
    parser.add_argument('--self-check', action='store_true', help='运行无网络合成数据自检后退出')
    args = parser.parse_args(argv)
    if not 1 <= args.port <= 65535:
        parser.error('port must be between 1 and 65535')
    try:
        if args.self_check:
            return self_check()
        import uvicorn
        from .app import create_app
        from .paths import default_db_path
        # Keep the socket reserved through startup, eliminating preflight/bind races.
        with reserve_socket(args.port) as sock:
            app = create_app()
            server = uvicorn.Server(uvicorn.Config(app, host='127.0.0.1', port=args.port,
                                    workers=1, access_log=False, loop='asyncio', http='h11', ws='none'))
            url = f'http://127.0.0.1:{args.port}'
            print(f'网页搬表：{url}\n数据保存在：{default_db_path()}\n停止服务请按 Ctrl+C；关闭网页不会停止服务。', flush=True)
            stop = threading.Event()
            browser = None
            if not args.no_browser:
                browser = threading.Thread(target=open_when_ready, args=(server, url, stop), daemon=True)
                browser.start()
            # Uvicorn re-raises captured signals after serving. Restore to our
            # cooperative handlers first, so SIGTERM/SIGBREAK also drain workers.
            handled = [signal.SIGINT, signal.SIGTERM]
            if hasattr(signal, 'SIGBREAK'):
                handled.append(signal.SIGBREAK)
            def request_stop(signum, frame):
                server.should_exit = True
            originals = {sig: signal.signal(sig, request_stop) for sig in handled}
            try:
                server.run(sockets=[sock])
            finally:
                stop.set()
                if browser:
                    browser.join(timeout=2)
                try:
                    app.state.store.shutdown()
                finally:
                    for sig, handler in originals.items():
                        signal.signal(sig, handler)
        return 0
    except KeyboardInterrupt:
        return 0
    except Exception as exc:
        # Never include keys, extracted page contents, or arbitrary exception bodies.
        message = ('启动失败：端口可能已被占用，或数据目录不可写。请关闭占用端口的程序，'
                   '并检查当前用户的数据目录权限后重试。不会打开或接管已有服务。'
                   f'\n错误类型：{type(exc).__name__}')
        print(message, file=sys.stderr, flush=True)
        if os.name == 'nt' and getattr(sys, 'frozen', False) and not args.no_browser and not args.self_check:
            import ctypes
            ctypes.windll.user32.MessageBoxW(None, message, '网页搬表启动失败', 0x10)
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
